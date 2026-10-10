import openpyxl
import re
import sys
import unicodedata
from collections import defaultdict
from difflib import SequenceMatcher

sys.stdout.reconfigure(encoding='utf-8')

LIB = r'E:\养娃乐\贝瓦儿歌\全新-贝瓦音乐.xlsx'
REQ = r'C:\Users\kings\Desktop\中文常见儿歌200首.xlsx'

def norm(s):
    s = '' if s is None else str(s)
    s = unicodedata.normalize('NFKC', s).lower()
    return re.sub(r'[\s\-_—()（）【】\[\]{}:：,，.!！？/、\'"·。]+', '', s)

def parts(s):
    s = str(s)
    s = re.sub(r'[（(].*?[）)]', '', s)
    return [norm(x) for x in re.split(r'[/／]', s) if norm(x)]

wl = openpyxl.load_workbook(LIB, read_only=True, data_only=True)['音频清单']
wr = openpyxl.load_workbook(REQ, read_only=True, data_only=True).active
lib_rows = list(wl.iter_rows(min_row=2, values_only=True))
idx = defaultdict(list)
for row in lib_rows:
    idx[norm(row[1])].append(row)

for r in wr.iter_rows(min_row=2, values_only=True):
    no, title, lang, classic = r
    candidates = []
    for p in parts(title):
        candidates += idx.get(p, [])
    seen = set()
    candidates = [x for x in candidates if x[0] not in seen and not seen.add(x[0])]
    if candidates:
        print(no, title, 'EXACT', len(candidates), '|', '; '.join(f'{x[0]}:{x[1]}[{x[4]}]' for x in candidates[:20]))

print('\nFUZZY TOP')
names = [(row, norm(row[1])) for row in lib_rows if row[1]]
for r in wr.iter_rows(min_row=2, values_only=True):
    no, title, lang, classic = r
    ps = parts(title)
    scored = []
    for row, nn in names:
        best = 0.0
        for p in ps:
            if not p or not nn:
                continue
            if p in nn or nn in p:
                score = min(len(p), len(nn)) / max(len(p), len(nn))
            else:
                score = SequenceMatcher(None, p, nn).ratio()
            best = max(best, score)
        scored.append((best, row))
    scored.sort(key=lambda x: (-x[0], str(x[1][1])))
    top = [(round(s, 3), row[0], row[1], row[4]) for s, row in scored[:5]]
    print(no, title, top)
